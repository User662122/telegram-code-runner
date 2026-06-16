"""
HIMSON TEXTILES ENGG IND PVT LTD — Full Export (single script)
===============================================================
Produces ONE Excel workbook with 4 sheets:
  Sheet 1 : Trial Balance          (full TB, all accounts)
  Sheet 2 : Active Accounts        (TB accounts that had FY transactions)
  Sheet 3 : Ledger Summary         (one row per ledger account)
  Sheet 4 : All Transactions       (every ledger entry, grouped by account)

Usage
-----
  pip install pdfplumber openpyxl
  python himson_full_export.py

Input PDFs
----------
  TB_PDF     : Trial Balance PDF
  LEDGER_PDF : Account Wise Ledger PDF

Output
------
  OUTPUT_XLSX : combined Excel workbook
"""

import pdfplumber, openpyxl, re, sys
from collections import defaultdict
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── File paths ───────────────────────────────────────────────────────────────
TB_PDF      = "/mnt/user-data/uploads/hkitbr1.pdf"
LEDGER_PDF  = "/mnt/user-data/uploads/hkiallldr.pdf"
OUTPUT_XLSX = "/mnt/user-data/outputs/himson_full_export.xlsx"

# ═══════════════════════════════════════════════════════════════════════════════
# SHARED HELPERS
# ═══════════════════════════════════════════════════════════════════════════════

def is_number(t):
    return bool(re.match(r'^[\d,]+\.\d+$', t.strip()))

def clean_num(t):
    try:    return float(t.replace(',', ''))
    except: return None

NF = '#,##0.00;(#,##0.00);"-"'

# Shared Excel styles (created once, reused across sheets)
HDR_FILL   = PatternFill("solid", start_color="1F4E79")
L0_FILL    = PatternFill("solid", start_color="BDD7EE")
L1_FILL    = PatternFill("solid", start_color="DDEBF7")
L2_FILL    = PatternFill("solid", start_color="EBF3FB")
TOTAL_FILL = PatternFill("solid", start_color="FCE4D6")
ALT_FILL   = PatternFill("solid", start_color="F5F9FF")
ACCT_FILL  = PatternFill("solid", start_color="2E75B6")
CL_FILL    = PatternFill("solid", start_color="E2EFDA")

HDR_FONT   = Font(name="Arial", bold=True, color="FFFFFF", size=10)
L0_FONT    = Font(name="Arial", bold=True, color="1F4E79", size=10)
L1_FONT    = Font(name="Arial", bold=True, color="17375E", size=9)
L2_FONT    = Font(name="Arial", bold=True, color="2E75B6", size=9)
DATA_FONT  = Font(name="Arial", size=9)
TOTAL_FONT = Font(name="Arial", bold=True, size=10, color="C00000")
ACCT_FONT  = Font(name="Arial", bold=True, color="FFFFFF", size=10)
CL_FONT    = Font(name="Arial", bold=True, size=9, color="375623")

def write_col_headers(ws, row, headers, widths, fill=HDR_FILL, font=HDR_FONT):
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = font; c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[row].height = 30

def group_header(ws, col_start, col_end, row, text, color):
    start_ltr = get_column_letter(col_start)
    end_ltr   = get_column_letter(col_end)
    ws.merge_cells(f"{start_ltr}{row}:{end_ltr}{row}")
    c = ws.cell(row=row, column=col_start, value=text)
    c.font = Font(name="Arial", bold=True, size=9, color="FFFFFF")
    c.fill = PatternFill("solid", start_color=color)
    c.alignment = Alignment(horizontal="center")

# ═══════════════════════════════════════════════════════════════════════════════
# PART 1 — TRIAL BALANCE PARSER
# ═══════════════════════════════════════════════════════════════════════════════

TB_COL_BOUNDS = [
    (0,   415, "ob_dr"),
    (415, 490, "ob_cr"),
    (490, 560, "tx_dr"),
    (560, 635, "tx_cr"),
    (635, 720, "cb_dr"),
    (720, 999, "cb_cr"),
]

TB_SKIP = {"Name", "Debit", "Credit", "Openning Balance", "Transaction",
           "Closing Balance", "Openning", "Balance", "Closing"}

def tb_assign_col(x1):
    for lo, hi, name in TB_COL_BOUNDS:
        if lo < x1 <= hi: return name
    return None

def tb_get_indent(x0):
    if x0 <= 40:  return 0
    if x0 <= 55:  return 1
    if x0 <= 68:  return 2
    return 3

def is_repeated_header_row(text_only):
    tokens = text_only.split()
    return bool(tokens) and all(t in ("Debit", "Credit") for t in tokens)

def parse_trial_balance():
    records = []
    with pdfplumber.open(TB_PDF) as pdf:
        for page in pdf.pages:
            words = page.extract_words(x_tolerance=3, y_tolerance=3)
            lines = defaultdict(list)
            for w in words:
                lines[round(w['top'])].append(w)
            sorted_ys = sorted(lines.keys())
            i = 0
            while i < len(sorted_ys):
                y = sorted_ys[i]
                line = sorted(lines[y], key=lambda w: w['x0'])
                text_only = " ".join(w['text'] for w in line)

                # Skip page banners / column headers
                if any(s in text_only for s in [
                    "HIMSON TEXTILES", "HIRALAL COLONY", "SURAT.", "Page :",
                    "DETAIL TRIAL", "01-04-25", "Openning", "Closing Balance",
                ]):
                    i += 1; continue
                if "Grand Total" in text_only:
                    nums = {tb_assign_col(round(w['x1'])): clean_num(w['text'])
                            for w in line if is_number(w['text'])}
                    if i + 1 < len(sorted_ys):
                        nxt = sorted(lines[sorted_ys[i + 1]], key=lambda w: w['x0'])
                        for w in nxt:
                            if is_number(w['text']):
                                col = tb_assign_col(round(w['x1']))
                                if col: nums[col] = clean_num(w['text'])
                    records.append({"name": "Grand Total", "indent": 0,
                                    **{k: nums.get(k) for k in ["ob_dr","ob_cr","tx_dr","tx_cr","cb_dr","cb_cr"]}})
                    i += 1; continue

                if all(is_number(w['text']) for w in line):
                    i += 1; continue
                if len(line) == 1 and line[0]['text'] in TB_SKIP:
                    i += 1; continue
                if is_repeated_header_row(text_only):
                    i += 1; continue

                name_words = [w for w in line if not is_number(w['text'])]
                num_words  = [w for w in line if is_number(w['text'])]
                if not name_words: i += 1; continue
                if all(w['top'] < 60 for w in name_words): i += 1; continue

                name   = " ".join(w['text'] for w in name_words)
                x0     = name_words[0]['x0']
                indent = tb_get_indent(x0)
                nums   = {}
                for w in num_words:
                    col = tb_assign_col(round(w['x1']))
                    if col: nums[col] = clean_num(w['text'])

                # Absorb a following all-number line as this row's values
                if i + 1 < len(sorted_ys):
                    nxt_y    = sorted_ys[i + 1]
                    nxt_line = sorted(lines[nxt_y], key=lambda w: w['x0'])
                    if nxt_line and all(is_number(w['text']) for w in nxt_line):
                        for w in nxt_line:
                            col = tb_assign_col(round(w['x1']))
                            if col: nums[col] = clean_num(w['text'])
                        i += 1

                records.append({"name": name, "indent": indent,
                                 **{k: nums.get(k) for k in ["ob_dr","ob_cr","tx_dr","tx_cr","cb_dr","cb_cr"]}})
                i += 1
    return records


def filter_active_tb(records):
    """Keep section headers + only leaf accounts that had FY transactions."""
    NUM_KEYS = ["ob_dr","ob_cr","tx_dr","tx_cr","cb_dr","cb_cr"]

    def has_nums(r):  return any(r.get(k) for k in NUM_KEYS)
    def has_txn(r):   return bool(r.get("tx_dr")) or bool(r.get("tx_cr"))

    tagged = []
    for r in records:
        if r["name"] == "Grand Total": tagged.append((r, "total"))
        elif not has_nums(r):          tagged.append((r, "header"))
        else:                          tagged.append((r, "leaf"))

    n    = len(tagged)
    keep = [False] * n

    # Mark active leaves and total
    for idx, (r, kind) in enumerate(tagged):
        if kind == "leaf" and has_txn(r):  keep[idx] = True
        elif kind == "total":              keep[idx] = True

    # Keep headers that have at least one active leaf / sub-header beneath them
    for idx, (r, kind) in enumerate(tagged):
        if kind != "header": continue
        j = idx + 1
        found = False
        while j < n:
            r2, kind2 = tagged[j]
            if kind2 == "header" and r2["indent"] <= r["indent"]: break
            if keep[j]: found = True; break
            j += 1
        keep[idx] = found

    return [r for (r, _), k in zip(tagged, keep) if k]


def write_tb_sheet(ws, records, title):
    ws.merge_cells("A1:G1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", bold=True, size=12, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center")

    group_header(ws, 2, 3, 2, "Opening Balance", "2E75B6")
    group_header(ws, 4, 5, 2, "Transaction",     "375623")
    group_header(ws, 6, 7, 2, "Closing Balance", "7030A0")

    write_col_headers(ws, 3,
        ["Account Name","Opening\nDebit","Opening\nCredit",
         "Transaction\nDebit","Transaction\nCredit",
         "Closing\nDebit","Closing\nCredit"],
        [55, 16, 16, 16, 16, 16, 16])

    ws.freeze_panes = "A4"
    alt = False; r = 4
    for rec in records:
        name      = rec["name"]
        indent    = rec["indent"]
        has_nums  = any(rec.get(k) for k in ["ob_dr","ob_cr","tx_dr","tx_cr","cb_dr","cb_cr"])
        is_total  = name == "Grand Total"

        if is_total:
            row_fill = TOTAL_FILL; name_font = row_font = TOTAL_FONT
        elif not has_nums and indent == 0:
            row_fill = L0_FILL; name_font = row_font = L0_FONT
        elif not has_nums and indent == 1:
            row_fill = L1_FILL; name_font = row_font = L1_FONT
        elif not has_nums and indent == 2:
            row_fill = L2_FILL; name_font = row_font = L2_FONT
        else:
            row_fill = ALT_FILL if alt else None
            name_font = row_font = DATA_FONT
            alt = not alt

        c = ws.cell(row=r, column=1, value=name)
        c.font = name_font
        c.alignment = Alignment(indent=indent, vertical="center")
        if row_fill: c.fill = row_fill

        for col_idx, key in enumerate(["ob_dr","ob_cr","tx_dr","tx_cr","cb_dr","cb_cr"], 2):
            cell = ws.cell(row=r, column=col_idx, value=rec.get(key))
            cell.font = row_font
            cell.number_format = NF
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if row_fill: cell.fill = row_fill

        if is_total:
            med = Side(style="medium", color="1F4E79")
            dbl = Side(style="double", color="1F4E79")
            for col_idx in range(1, 8):
                ws.cell(row=r, column=col_idx).border = Border(top=med, bottom=dbl)
        r += 1

    ws.auto_filter.ref = f"A3:{get_column_letter(7)}{r - 1}"
    return r


# ═══════════════════════════════════════════════════════════════════════════════
# PART 2 — LEDGER PARSER  (fixed)
# ═══════════════════════════════════════════════════════════════════════════════

def is_date6(t):  return bool(re.fullmatch(r'\d{2}/\d{2}/\d{2}', t))
def is_num_l(t):  return bool(re.fullmatch(r'[\d,]+\.?\d*', t))
def to_f(t):
    try:    return float(t.replace(',', ''))
    except: return None

def ledger_num_col(x1):
    if   410 < x1 <= 462: return 'debit'
    elif 462 < x1 <= 522: return 'credit'
    elif 522 < x1 <= 585: return 'balance'
    return None

LEDGER_SKIP_HDR = {
    'HIMSON','TEXTILES','ENGG','IND','PVT','LTD','HIRALAL',
    'COLONY,,A.K.ROAD,','SURAT.','Account','Wise','Ledger',
    'From','To','Vch.','Ref.','No.','Date','Narration',
    'Debit','Credit','Balance','Page','#'
}

LEDGER_SKIP_KW = {
    'TRANSECTION','TRANSACTION','TOTAL','CLOSING','BALANCE',
    'BEING','Op.','Op-Bal','Op'
}

def ledger_page_lines(page):
    words = page.extract_words(x_tolerance=3, y_tolerance=3)
    d = defaultdict(list)
    for w in words:
        d[round(w['top'])].append(w)
    result = []
    for y in sorted(d):
        if y < 65: continue
        line = sorted(d[y], key=lambda w: w['x0'])
        if all(w['text'] in LEDGER_SKIP_HDR for w in line): continue
        result.append((y, line))
    return result


def parse_ledger():
    accounts  = []
    cur       = None
    cur_txn   = None
    narr_buf  = []
    pending_dc       = None
    expect_cl_bal    = False
    expect_cl_dc     = False   # next lone D/C after CLOSING BALANCE is for the account
    continuation_op  = False   # True when we've just seen a page-break op-balance line

    def push_txn():
        nonlocal cur_txn, narr_buf
        if cur_txn and cur:
            cur_txn['narration'] = ' '.join(narr_buf).strip()
            cur['transactions'].append(cur_txn)
        cur_txn = None
        narr_buf = []

    def push_acct():
        nonlocal cur
        push_txn()
        if cur: accounts.append(cur)
        cur = None

    def new_acct(name):
        nonlocal cur
        push_acct()
        cur = dict(name=name, op_bal=None, op_dc=None,
                   transactions=[], tot_dr=None, tot_cr=None,
                   cl_bal=None, cl_dc=None)

    with pdfplumber.open(LEDGER_PDF) as pdf:
        total = len(pdf.pages)
        for pn, page in enumerate(pdf.pages):
            if pn % 50 == 0:
                print(f"  Ledger page {pn+1}/{total}", flush=True)

            lines = ledger_page_lines(page)
            i = 0
            while i < len(lines):
                y, line = lines[i]
                texts  = [w['text'] for w in line]
                joined = ' '.join(texts)

                # ── Closing balance number (line immediately after CLOSING BALANCE label) ──
                if expect_cl_bal:
                    expect_cl_bal = False
                    nums = [w for w in line if is_num_l(w['text'])]
                    if nums and cur:
                        cur['cl_bal'] = to_f(nums[0]['text'])
                    # DC may follow on same line or on next standalone D/C line
                    dc_tok = next((w['text'] for w in line
                                   if w['text'] in ('D','C') and w['x0'] > 300), None)
                    if dc_tok and cur:
                        cur['cl_dc'] = dc_tok
                    elif cur:
                        expect_cl_dc = True   # watch for standalone D/C next
                    i += 1; continue

                # ── Standalone D / C flag ─────────────────────────────────
                if len(line) == 1 and texts[0] in ('D', 'C') and line[0]['x0'] > 300:
                    pending_dc = texts[0]
                    if expect_cl_dc and cur:
                        cur['cl_dc'] = texts[0]
                        expect_cl_dc = False
                    elif cur_txn and cur_txn['dc'] is None:
                        cur_txn['dc'] = texts[0]
                    i += 1; continue

                expect_cl_dc = False   # reset if we hit any other line

                # ── Op. Balance ───────────────────────────────────────────
                if 'Op.' in texts and 'Balance' in texts:
                    nums = [w for w in line if is_num_l(w['text'])]
                    dc   = next((w['text'] for w in line
                                 if w['text'] in ('D','C') and w['x0'] > 300), None)
                    if not nums and i + 1 < len(lines):
                        _, nxt = lines[i + 1]
                        nums_nxt = [w for w in nxt if is_num_l(w['text'])]
                        if nums_nxt:
                            nums = nums_nxt
                            dc = dc or next((w['text'] for w in nxt
                                            if w['text'] in ('D','C') and w['x0'] > 300), None)
                            i += 1
                    if cur:
                        # Only set op_bal if this is the FIRST occurrence (not a page-break re-statement)
                        if cur['op_bal'] is None:
                            cur['op_bal'] = to_f(nums[0]['text']) if nums else None
                            cur['op_dc']  = dc or pending_dc
                        # If op_bal already set, this is a page-break re-statement — ignore
                    pending_dc = None
                    i += 1; continue

                # ── TRANSACTION TOTAL ─────────────────────────────────────
                if re.search(r'TRANS[EA]CTION\s+TOTAL', joined):
                    push_txn()
                    nums = {}
                    for w in line:
                        if is_num_l(w['text']):
                            c = ledger_num_col(round(w['x1']))
                            if c: nums[c] = to_f(w['text'])
                    if cur:
                        cur['tot_dr'] = nums.get('debit')
                        cur['tot_cr'] = nums.get('credit')
                    pending_dc = None
                    i += 1; continue

                # ── TOTAL (running sub-total, skip) ───────────────────────
                if texts and texts[0] == 'TOTAL':
                    i += 1; continue

                # ── CLOSING BALANCE label ─────────────────────────────────
                if 'CLOSING' in texts and 'BALANCE' in texts:
                    push_txn()
                    expect_cl_bal = True
                    pending_dc    = None
                    i += 1; continue

                # ── Transaction date line (date in col 0) ─────────────────
                if texts and is_date6(texts[0]) and line[0]['x0'] < 50:
                    push_txn()
                    dc_flag = pending_dc; pending_dc = None
                    date = texts[0]
                    vch_type = vch_no = ref_no = ref_date = ''

                    for w in line:
                        x0 = round(w['x0']); t = w['text']
                        if is_date6(t) and x0 < 50:          date = t
                        elif 55 <= x0 < 88  and not is_date6(t) and not is_num_l(t):
                            vch_type = (vch_type + ' ' + t).strip()
                        elif 88 <= x0 < 130 and not is_date6(t):
                            vch_no   = (vch_no   + ' ' + t).strip()
                        elif 130 <= x0 < 185 and not is_date6(t):
                            ref_no   = (ref_no   + ' ' + t).strip()
                        elif 185 <= x0 < 230 and is_date6(t):
                            ref_date = t

                    # Sometimes vch_type wraps to next line
                    if not vch_type and i + 1 < len(lines):
                        _, nxt = lines[i + 1]
                        maybe = [w for w in nxt
                                 if 55 <= w['x0'] < 88
                                 and not is_date6(w['text'])
                                 and not is_num_l(w['text'])]
                        if maybe:
                            for w in nxt:
                                x0 = round(w['x0']); t = w['text']
                                if 55 <= x0 < 88  and not is_date6(t) and not is_num_l(t):
                                    vch_type = (vch_type + ' ' + t).strip()
                                elif 88 <= x0 < 130 and not is_date6(t):
                                    vch_no   = (vch_no   + ' ' + t).strip()
                                elif 130 <= x0 < 185 and not is_date6(t):
                                    ref_no   = (ref_no   + ' ' + t).strip()
                                elif 185 <= x0 < 230 and is_date6(t):
                                    ref_date = t
                            i += 1

                    cur_txn = dict(date=date, vch_type=vch_type, vch_no=vch_no,
                                   ref_no=ref_no, ref_date=ref_date,
                                   debit=None, credit=None, balance=None, dc=dc_flag)
                    narr_buf = []
                    i += 1; continue

                # ── Narration / amount lines (x0 >= 55, not a date) ───────
                if line and line[0]['x0'] >= 55 and not is_date6(texts[0]):
                    num_ws  = [w for w in line if is_num_l(w['text'])]
                    text_ws = [w for w in line
                               if not is_num_l(w['text']) and w['text'] not in ('D','C')]
                    dc_tok  = next((w['text'] for w in line
                                    if w['text'] in ('D','C') and w['x0'] > 560), None)
                    if cur_txn:
                        for w in num_ws:
                            col = ledger_num_col(round(w['x1']))
                            if col == 'debit'   and cur_txn['debit']   is None:
                                cur_txn['debit']   = to_f(w['text'])
                            elif col == 'credit' and cur_txn['credit']  is None:
                                cur_txn['credit']  = to_f(w['text'])
                            elif col == 'balance' and cur_txn['balance'] is None:
                                cur_txn['balance'] = to_f(w['text'])
                        if dc_tok: cur_txn['dc'] = dc_tok
                        if text_ws: narr_buf.extend(w['text'] for w in text_ws)
                    i += 1; continue

                # ── Account name line (x0 <= 30, not a date) ─────────────
                if line and line[0]['x0'] <= 30 and not is_date6(texts[0]):
                    name_ws = [w for w in line
                               if w['x0'] <= 420 and w['text'] not in ('D','C')]
                    name    = ' '.join(w['text'] for w in name_ws).strip()

                    if name and not any(k in name for k in LEDGER_SKIP_KW):
                        push_txn()

                        # ── FIX: correct page-continuation detection ──────
                        # A page break re-prints the current account's name.
                        # Detect: same name as active account AND no closing
                        # balance yet (account is still open).
                        is_continuation = (
                            cur is not None
                            and cur['cl_bal'] is None
                            and cur['name'] == name
                        )
                        if not is_continuation:
                            new_acct(name)
                        # If it IS a continuation we simply stay on cur

                    pending_dc = None
                    i += 1; continue

                i += 1

    push_acct()
    return accounts


def merge_ledger_accounts(raw):
    """
    Merge page-split segments of the same account.
    A segment is a continuation when:
      - Same name as previous
      - Previous segment has no closing balance yet
    We keep the FIRST segment's op_bal/op_dc and take the LAST segment's
    tot_dr/tot_cr/cl_bal/cl_dc.
    """
    merged = []
    for a in raw:
        if (merged
                and merged[-1]['name'] == a['name']
                and merged[-1]['cl_bal'] is None):
            prev = merged[-1]
            prev['transactions'].extend(a['transactions'])
            # Always take later segment's totals / closing (they're cumulative)
            if a['tot_dr']  is not None: prev['tot_dr']  = a['tot_dr']
            if a['tot_cr']  is not None: prev['tot_cr']  = a['tot_cr']
            if a['cl_bal']  is not None: prev['cl_bal']  = a['cl_bal']
            if a['cl_dc']   is not None: prev['cl_dc']   = a['cl_dc']
            # Do NOT overwrite op_bal — keep the original opening balance
        else:
            merged.append(a)
    return merged


# ═══════════════════════════════════════════════════════════════════════════════
# PART 3 — WRITE LEDGER SHEETS
# ═══════════════════════════════════════════════════════════════════════════════

def write_ledger_summary(ws, accounts):
    ws.merge_cells("A1:H1")
    ws["A1"] = "HIMSON TEXTILES ENGG IND PVT LTD — Ledger Summary (01-Apr-2025 to 31-Mar-2026)"
    ws["A1"].font = Font(name="Arial", bold=True, size=12, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center")

    write_col_headers(ws, 3,
        ["Account Name", "Opening Balance", "Dr/Cr", "Total Debit",
         "Total Credit", "Closing Balance", "Dr/Cr", "No. of Txns"],
        [52, 18, 7, 18, 18, 18, 7, 12])

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = "A3:H3"

    sr = 4; alt = False
    for a in accounts:
        f = ALT_FILL if alt else None; alt = not alt
        for col, val in enumerate([
            a['name'], a['op_bal'], a['op_dc'], a['tot_dr'],
            a['tot_cr'], a['cl_bal'], a['cl_dc'], len(a['transactions'])
        ], 1):
            c = ws.cell(row=sr, column=col, value=val)
            c.font = DATA_FONT
            if f: c.fill = f
            if col in (2, 4, 5, 6):
                c.number_format = NF
                c.alignment = Alignment(horizontal="right")
            elif col in (3, 7, 8):
                c.alignment = Alignment(horizontal="center")
        sr += 1

    # Grand total row
    for col in range(1, 9):
        c = ws.cell(row=sr, column=col); c.fill = TOTAL_FILL; c.font = TOTAL_FONT
    ws.cell(row=sr, column=1).value = "GRAND TOTAL"
    for col, ltr in [(2, "B"), (4, "D"), (5, "E"), (6, "F")]:
        c = ws.cell(row=sr, column=col)
        c.value = f"=SUM({ltr}4:{ltr}{sr-1})"
        c.number_format = NF
        c.font = TOTAL_FONT; c.fill = TOTAL_FILL
        c.alignment = Alignment(horizontal="right")
    c = ws.cell(row=sr, column=8)
    c.value = f"=SUM(H4:H{sr-1})"
    c.font = TOTAL_FONT; c.fill = TOTAL_FILL
    c.alignment = Alignment(horizontal="center")

    return sr


def write_all_transactions(ws, accounts):
    ws.merge_cells("A1:K1")
    ws["A1"] = "HIMSON TEXTILES ENGG IND PVT LTD — All Ledger Transactions (01-Apr-2025 to 31-Mar-2026)"
    ws["A1"].font = Font(name="Arial", bold=True, size=12, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center")

    write_col_headers(ws, 3,
        ["Account Name", "Date", "Vch Type", "Vch No", "Ref No",
         "Ref Date", "Narration", "Debit", "Credit", "Balance", "Dr/Cr"],
        [40, 11, 9, 9, 12, 11, 55, 15, 15, 15, 6])

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(11)}3"

    dr = 4; alt_d = False
    for a in accounts:
        # Account header row
        ws.merge_cells(f"A{dr}:F{dr}")
        c = ws.cell(row=dr, column=1, value=a['name'])
        c.font = ACCT_FONT; c.fill = ACCT_FILL
        c.alignment = Alignment(vertical="center")
        ob = (f"Op. Bal: {a['op_bal']:,.2f} {a['op_dc'] or ''}"
              if a['op_bal'] is not None else "Op. Bal: —")
        c2 = ws.cell(row=dr, column=7, value=ob)
        c2.font = Font(name="Arial", bold=True, color="FFFFFF", size=9, italic=True)
        c2.fill = ACCT_FILL
        for col in range(2, 12): ws.cell(row=dr, column=col).fill = ACCT_FILL
        dr += 1

        for txn in a['transactions']:
            f = ALT_FILL if alt_d else None; alt_d = not alt_d
            for col, val in enumerate([
                a['name'], txn['date'], txn['vch_type'], txn['vch_no'],
                txn['ref_no'], txn['ref_date'], txn['narration'],
                txn['debit'], txn['credit'], txn['balance'], txn['dc']
            ], 1):
                c = ws.cell(row=dr, column=col, value=val)
                c.font = DATA_FONT
                if f: c.fill = f
                if col in (8, 9, 10):
                    c.number_format = NF
                    c.alignment = Alignment(horizontal="right", vertical="center")
                elif col == 11:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif col == 7:
                    c.alignment = Alignment(wrap_text=True, vertical="center")
                else:
                    c.alignment = Alignment(vertical="center")
            dr += 1

        # Closing balance footer row
        ws.merge_cells(f"A{dr}:G{dr}")
        cb = (f"Closing Bal: {a['cl_bal']:,.2f} {a['cl_dc'] or ''}"
              if a['cl_bal'] is not None else "Closing Bal: —")
        c = ws.cell(row=dr, column=1, value=cb)
        c.font = CL_FONT; c.fill = CL_FILL
        c.alignment = Alignment(horizontal="right")
        for col in range(2, 12): ws.cell(row=dr, column=col).fill = CL_FILL
        dr += 2   # blank gap between accounts

    return dr


# ═══════════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════════

def main():
    wb = openpyxl.Workbook()

    # ── Sheet 1: Full Trial Balance ──────────────────────────────────────────
    print("\n[1/4] Parsing Trial Balance PDF…")
    tb_records = parse_trial_balance()
    print(f"      Parsed {len(tb_records)} rows")

    ws_tb = wb.active
    ws_tb.title = "Trial Balance"
    write_tb_sheet(ws_tb, tb_records,
                   "HIMSON TEXTILES ENGG IND PVT LTD — DETAIL TRIAL BALANCE (01-Apr-2025 to 31-Mar-2026)")
    print(f"      Sheet 'Trial Balance' written")

    # ── Sheet 2: Active Accounts ─────────────────────────────────────────────
    print("\n[2/4] Filtering active accounts (had FY transactions)…")
    active_records = filter_active_tb(tb_records)
    leaf_count = sum(
        1 for r in active_records
        if any(r.get(k) for k in ["ob_dr","ob_cr","tx_dr","tx_cr","cb_dr","cb_cr"])
        and r["name"] != "Grand Total"
    )
    print(f"      {len(active_records)} rows kept ({leaf_count} active leaf accounts)")

    ws_active = wb.create_sheet("Active Accounts (FY Txns)")
    ws_active.merge_cells("A1:G1")
    ws_active["A1"] = "HIMSON TEXTILES ENGG IND PVT LTD — ACCOUNTS WITH FY TRANSACTIONS (01-Apr-2025 to 31-Mar-2026)"
    ws_active["A1"].font = Font(name="Arial", bold=True, size=12, color="1F4E79")
    ws_active["A1"].alignment = Alignment(horizontal="center")
    ws_active.merge_cells("A2:G2")
    ws_active["A2"] = "Accounts with zero transaction movement during the year are excluded"
    ws_active["A2"].font = Font(name="Arial", italic=True, size=9, color="595959")
    ws_active["A2"].alignment = Alignment(horizontal="center")
    group_header(ws_active, 2, 3, 3, "Opening Balance", "2E75B6")
    group_header(ws_active, 4, 5, 3, "Transaction",     "375623")
    group_header(ws_active, 6, 7, 3, "Closing Balance", "7030A0")
    write_col_headers(ws_active, 4,
        ["Account Name","Opening\nDebit","Opening\nCredit",
         "Transaction\nDebit","Transaction\nCredit",
         "Closing\nDebit","Closing\nCredit"],
        [55, 16, 16, 16, 16, 16, 16])
    ws_active.freeze_panes = "A5"

    alt = False; r = 5
    for rec in active_records:
        name     = rec["name"]
        indent   = rec["indent"]
        has_nums = any(rec.get(k) for k in ["ob_dr","ob_cr","tx_dr","tx_cr","cb_dr","cb_cr"])
        is_total = name == "Grand Total"

        if is_total:
            row_fill = TOTAL_FILL; name_font = row_font = TOTAL_FONT
        elif not has_nums and indent == 0:
            row_fill = L0_FILL; name_font = row_font = L0_FONT
        elif not has_nums and indent == 1:
            row_fill = L1_FILL; name_font = row_font = L1_FONT
        elif not has_nums and indent == 2:
            row_fill = L2_FILL; name_font = row_font = L2_FONT
        else:
            row_fill = ALT_FILL if alt else None
            name_font = row_font = DATA_FONT
            alt = not alt

        c = ws_active.cell(row=r, column=1, value=name)
        c.font = name_font
        c.alignment = Alignment(indent=indent, vertical="center")
        if row_fill: c.fill = row_fill

        for col_idx, key in enumerate(["ob_dr","ob_cr","tx_dr","tx_cr","cb_dr","cb_cr"], 2):
            cell = ws_active.cell(row=r, column=col_idx, value=rec.get(key))
            cell.font = row_font; cell.number_format = NF
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if row_fill: cell.fill = row_fill

        if is_total:
            med = Side(style="medium", color="1F4E79")
            dbl = Side(style="double", color="1F4E79")
            for col_idx in range(1, 8):
                ws_active.cell(row=r, column=col_idx).border = Border(top=med, bottom=dbl)
        r += 1

    ws_active.auto_filter.ref = f"A4:{get_column_letter(7)}{r - 1}"
    print(f"      Sheet 'Active Accounts (FY Txns)' written")

    # ── Sheets 3 & 4: Ledger ─────────────────────────────────────────────────
    print("\n[3/4] Parsing Ledger PDF…")
    try:
        raw_ledger = parse_ledger()
        print(f"      Parsed {len(raw_ledger)} raw account segments")

        ledger_accounts = merge_ledger_accounts(raw_ledger)
        print(f"      Merged into {len(ledger_accounts)} accounts, "
              f"{sum(len(a['transactions']) for a in ledger_accounts)} transactions")

        # Cross-check: compare ledger account names against TB leaf names
        tb_leaf_names = {
            r["name"] for r in tb_records
            if any(r.get(k) for k in ["ob_dr","ob_cr","tx_dr","tx_cr","cb_dr","cb_cr"])
            and r["name"] != "Grand Total"
        }
        ledger_names = {a["name"] for a in ledger_accounts}
        only_in_ledger = ledger_names - tb_leaf_names
        only_in_tb     = tb_leaf_names - ledger_names
        if only_in_ledger:
            print(f"\n  ⚠  {len(only_in_ledger)} ledger accounts NOT found in TB (possible duplicates/parse errors):")
            for n in sorted(only_in_ledger): print(f"       - {n}")
        if only_in_tb:
            print(f"\n  ⚠  {len(only_in_tb)} TB accounts NOT found in ledger:")
            for n in sorted(only_in_tb): print(f"       - {n}")

        print("\n[4/4] Writing ledger sheets…")
        ws_sum = wb.create_sheet("Ledger Summary")
        write_ledger_summary(ws_sum, ledger_accounts)
        print(f"      Sheet 'Ledger Summary' written ({len(ledger_accounts)} accounts)")

        ws_txn = wb.create_sheet("All Transactions")
        write_all_transactions(ws_txn, ledger_accounts)
        print(f"      Sheet 'All Transactions' written")

    except FileNotFoundError:
        print(f"  ⚠  Ledger PDF not found at: {LEDGER_PDF}")
        print("     Skipping ledger sheets — Trial Balance sheets still saved.")

    # ── Save ─────────────────────────────────────────────────────────────────
    wb.save(OUTPUT_XLSX)
    print(f"\n✓ Saved: {OUTPUT_XLSX}")
    print(f"  Sheets: {', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()
